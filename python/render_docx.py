"""Render approved structured JSON into DOCX and write a provenance manifest."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import ZipFile


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_payload(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("Input JSON must contain an object at the top level.")
    return payload


def render_without_template(payload: dict[str, Any], output_path: Path) -> None:
    from docx import Document

    title = payload.get("title")
    sections = payload.get("sections", [])
    if title is not None and not isinstance(title, str):
        raise ValueError("title must be a string when supplied.")
    if not isinstance(sections, list):
        raise ValueError("sections must be an array.")

    document = Document()
    if title:
        document.add_heading(title, level=0)

    for index, section in enumerate(sections):
        if not isinstance(section, dict):
            raise ValueError(f"sections[{index}] must be an object.")
        heading = section.get("heading")
        paragraphs = section.get("paragraphs", [])
        if heading is not None and not isinstance(heading, str):
            raise ValueError(f"sections[{index}].heading must be a string.")
        if isinstance(paragraphs, str):
            paragraphs = [paragraphs]
        if not isinstance(paragraphs, list) or not all(isinstance(item, str) for item in paragraphs):
            raise ValueError(f"sections[{index}].paragraphs must be a string or an array of strings.")
        if heading:
            document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)

    document.save(output_path)


def render_with_template(payload: dict[str, Any], template_path: Path, output_path: Path) -> None:
    from docxtpl import DocxTemplate

    template = DocxTemplate(template_path)
    template.render(payload)
    template.save(output_path)


def validate_docx(path: Path) -> None:
    with ZipFile(path) as archive:
        names = set(archive.namelist())
    required = {"[Content_Types].xml", "word/document.xml"}
    missing = required - names
    if missing:
        raise ValueError(f"Generated DOCX is missing required entries: {', '.join(sorted(missing))}.")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_json", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument("--template", type=Path)
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path: Path = args.input_json.resolve()
    output_path: Path = args.output_docx.resolve()
    template_path: Path | None = args.template.resolve() if args.template else None
    manifest_path: Path = (
        args.manifest.resolve()
        if args.manifest
        else output_path.with_suffix(f"{output_path.suffix}.manifest.json")
    )

    if output_path.suffix.lower() != ".docx":
        raise ValueError("output_docx must use the .docx extension.")
    if output_path.exists() and not args.overwrite:
        raise FileExistsError(f"Output already exists: {output_path}")
    if manifest_path.exists() and not args.overwrite:
        raise FileExistsError(f"Manifest already exists: {manifest_path}")
    if template_path is not None and template_path.suffix.lower() != ".docx":
        raise ValueError("--template must reference a .docx file.")

    payload = load_payload(input_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if template_path is None:
        render_without_template(payload, output_path)
    else:
        render_with_template(payload, template_path, output_path)
    validate_docx(output_path)

    manifest = {
        "schemaVersion": 1,
        "tool": "render_docx.py",
        "createdAt": datetime.now(timezone.utc).isoformat(),
        "input": {
            "path": str(input_path),
            "sha256": sha256_file(input_path),
        },
        "template": None
        if template_path is None
        else {
            "path": str(template_path),
            "sha256": sha256_file(template_path),
        },
        "output": {
            "path": str(output_path),
            "sha256": sha256_file(output_path),
            "bytes": output_path.stat().st_size,
        },
        "reviewStatus": "unreviewed",
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(f"{json.dumps(manifest, indent=2)}\n", encoding="utf-8")
    print(json.dumps(manifest))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
